from __future__ import annotations

import csv
import json
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from .config import DB_PATH, RESULTS_DIR, ensure_storage


def utcnow() -> str:
    return datetime.utcnow().isoformat(timespec="seconds") + "Z"


def new_id() -> str:
    return str(uuid.uuid4())


def connect() -> sqlite3.Connection:
    ensure_storage()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS jobs (
                id TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                status TEXT NOT NULL,
                total_files INTEGER NOT NULL DEFAULT 0,
                processed_files INTEGER NOT NULL DEFAULT 0,
                current_file TEXT,
                progress INTEGER NOT NULL DEFAULT 0,
                message TEXT
            );

            CREATE TABLE IF NOT EXISTS studies (
                id TEXT PRIMARY KEY,
                job_id TEXT,
                created_at TEXT NOT NULL,
                display_name TEXT NOT NULL,
                group_name TEXT NOT NULL,
                source_type TEXT NOT NULL,
                source_path TEXT NOT NULL,
                status TEXT NOT NULL,
                anatomical_region TEXT,
                quality_class INTEGER,
                quality_prob REAL,
                violation_type TEXT,
                processing_status TEXT,
                time_of_processing REAL,
                preview_path TEXT,
                overlay_path TEXT,
                contour_path TEXT,
                metadata_json TEXT NOT NULL DEFAULT '{}',
                FOREIGN KEY(job_id) REFERENCES jobs(id)
            );

            CREATE TABLE IF NOT EXISTS contours (
                id TEXT PRIMARY KEY,
                study_id TEXT NOT NULL,
                version INTEGER NOT NULL,
                source TEXT NOT NULL,
                payload_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(study_id) REFERENCES studies(id)
            );
            """
        )


def row_to_dict(row: sqlite3.Row | None) -> dict[str, Any] | None:
    if row is None:
        return None
    data = dict(row)
    for key in ("metadata_json",):
        if key in data:
            try:
                data[key.removesuffix("_json")] = json.loads(data[key] or "{}")
            except json.JSONDecodeError:
                data[key.removesuffix("_json")] = {}
    return data


def rows_to_dicts(rows: list[sqlite3.Row]) -> list[dict[str, Any]]:
    return [row_to_dict(row) or {} for row in rows]


def create_job(total_files: int) -> str:
    job_id = new_id()
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO jobs (id, created_at, status, total_files, message)
            VALUES (?, ?, 'uploaded', ?, ?)
            """,
            (job_id, utcnow(), total_files, "Файлы загружены и ожидают обработки"),
        )
    return job_id


def add_study(
    *,
    job_id: str,
    display_name: str,
    group_name: str,
    source_type: str,
    source_path: Path,
    metadata: dict[str, Any] | None = None,
) -> str:
    study_id = new_id()
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO studies (
                id, job_id, created_at, display_name, group_name, source_type,
                source_path, status, metadata_json
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, 'uploaded', ?)
            """,
            (
                study_id,
                job_id,
                utcnow(),
                display_name,
                group_name,
                source_type,
                str(source_path),
                json.dumps(metadata or {}, ensure_ascii=False),
            ),
        )
    return study_id


def get_job(job_id: str) -> dict[str, Any] | None:
    with connect() as conn:
        row = conn.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
    return row_to_dict(row)


def get_study(study_id: str) -> dict[str, Any] | None:
    with connect() as conn:
        row = conn.execute("SELECT * FROM studies WHERE id = ?", (study_id,)).fetchone()
    return row_to_dict(row)


def list_studies(job_id: str | None = None) -> list[dict[str, Any]]:
    query = "SELECT * FROM studies"
    params: tuple[Any, ...] = ()
    if job_id:
        query += " WHERE job_id = ?"
        params = (job_id,)
    query += " ORDER BY created_at DESC, group_name ASC, display_name ASC"
    with connect() as conn:
        rows = conn.execute(query, params).fetchall()
    return rows_to_dicts(rows)


def list_jobs() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute("SELECT * FROM jobs ORDER BY created_at DESC").fetchall()
    return rows_to_dicts(rows)


def update_job(job_id: str, **fields: Any) -> None:
    if not fields:
        return
    keys = ", ".join(f"{key} = ?" for key in fields)
    values = tuple(fields.values()) + (job_id,)
    with connect() as conn:
        conn.execute(f"UPDATE jobs SET {keys} WHERE id = ?", values)


def update_study(study_id: str, **fields: Any) -> None:
    if not fields:
        return
    if "metadata" in fields:
        fields["metadata_json"] = json.dumps(fields.pop("metadata"), ensure_ascii=False)
    keys = ", ".join(f"{key} = ?" for key in fields)
    values = tuple(fields.values()) + (study_id,)
    with connect() as conn:
        conn.execute(f"UPDATE studies SET {keys} WHERE id = ?", values)


def delete_job(job_id: str) -> None:
    with connect() as conn:
        study_ids = [
            row[0]
            for row in conn.execute("SELECT id FROM studies WHERE job_id = ?", (job_id,)).fetchall()
        ]
        if study_ids:
            placeholders = ", ".join("?" for _ in study_ids)
            conn.execute(f"DELETE FROM contours WHERE study_id IN ({placeholders})", tuple(study_ids))
        conn.execute("DELETE FROM studies WHERE job_id = ?", (job_id,))
        conn.execute("DELETE FROM jobs WHERE id = ?", (job_id,))


def add_contour(study_id: str, payload: dict[str, Any], source: str = "ai") -> str:
    with connect() as conn:
        version = conn.execute(
            "SELECT COALESCE(MAX(version), 0) + 1 FROM contours WHERE study_id = ?",
            (study_id,),
        ).fetchone()[0]
        contour_id = new_id()
        conn.execute(
            """
            INSERT INTO contours (id, study_id, version, source, payload_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                contour_id,
                study_id,
                version,
                source,
                json.dumps(payload, ensure_ascii=False),
                utcnow(),
            ),
        )
    return contour_id


def list_contours(study_id: str) -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            "SELECT * FROM contours WHERE study_id = ? ORDER BY version DESC",
            (study_id,),
        ).fetchall()
    contours = []
    for row in rows:
        item = dict(row)
        item["payload"] = json.loads(item.pop("payload_json") or "{}")
        contours.append(item)
    return contours


def get_contour(contour_id: str) -> dict[str, Any] | None:
    with connect() as conn:
        row = conn.execute("SELECT * FROM contours WHERE id = ?", (contour_id,)).fetchone()
    if row is None:
        return None
    item = dict(row)
    item["payload"] = json.loads(item.pop("payload_json") or "{}")
    return item


def delete_contour(contour_id: str) -> dict[str, Any] | None:
    contour = get_contour(contour_id)
    if contour is None:
        return None
    with connect() as conn:
        conn.execute("DELETE FROM contours WHERE id = ?", (contour_id,))
    remaining = list_contours(contour["study_id"])
    manual_remaining = [
        item
        for item in remaining
        if item.get("source") == "manual" or (item.get("payload") or {}).get("source") == "manual_edit"
    ]
    update_study(
        contour["study_id"],
        contour_path=manual_remaining[0]["id"] if manual_remaining else None,
    )
    return contour


def write_studies_report(studies: list[dict[str, Any]], output_path: Path, output_format: str) -> Path:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    columns = [
        "path_to_study",
        "study_uid",
        "image_uid",
        "anatomical_region",
        "quality_class",
        "violation_type",
        "processing_status",
        "time_of_processing",
    ]

    rows = []
    for study in studies:
        metadata = study.get("metadata") or {}
        violation_type = study.get("violation_type") or ""
        if (study.get("processing_status") == "Failure") and not violation_type:
            violation_type = metadata.get("error_message") or metadata.get("validation_message", "")
        rows.append(
            {
                "path_to_study": study["source_path"],
                "study_uid": metadata.get("study_uid", ""),
                "image_uid": metadata.get("image_uid", ""),
                "anatomical_region": study.get("anatomical_region") or "",
                "quality_class": study.get("quality_class")
                if study.get("quality_class") is not None
                else "",
                "violation_type": violation_type,
                "processing_status": study.get("processing_status") or "",
                "time_of_processing": study.get("time_of_processing") or "",
            }
        )

    if output_format == "csv":
        with output_path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.DictWriter(file, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)
        return output_path

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "results"
    sheet.append(columns)
    for row in rows:
        sheet.append([row[column] for column in columns])
    workbook.save(output_path)
    return output_path


def export_results(job_id: str, output_format: str) -> Path:
    studies = list_studies(job_id)
    output_path = RESULTS_DIR / f"{job_id}.{output_format}"
    return write_studies_report(studies, output_path, output_format)


def export_selected_studies(study_ids: list[str], output_format: str) -> Path:
    if not study_ids:
        raise ValueError("no_studies_selected")
    placeholders = ", ".join("?" for _ in study_ids)
    with connect() as conn:
        rows = conn.execute(
            f"SELECT * FROM studies WHERE id IN ({placeholders})",
            tuple(study_ids),
        ).fetchall()
    by_id = {item["id"]: item for item in rows_to_dicts(rows)}
    studies = [by_id[study_id] for study_id in study_ids if study_id in by_id]
    output_path = RESULTS_DIR / f"archive-selected-{utcnow().replace(':', '').replace('-', '')}.{output_format}"
    return write_studies_report(studies, output_path, output_format)
